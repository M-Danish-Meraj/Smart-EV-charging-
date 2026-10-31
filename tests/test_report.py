"""Step 10: the report builder turns results.json into a complete .docx and .html,
with every key number taken from the results (none typed by hand)."""

import json
import subprocess
import sys
from pathlib import Path

import pytest
from docx import Document

ROOT = Path(__file__).resolve().parents[1]


@pytest.fixture(scope="module")
def built(tmp_path_factory):
    import run_all
    out = tmp_path_factory.mktemp("output")
    run_all.main(out, quick=True)
    rep = out / "report"
    subprocess.run([sys.executable, str(ROOT / "scripts" / "build_report.py"),
                    "--results", str(out / "results.json"), "--figures", str(out / "figures"),
                    "--out", str(rep), "--no-pdf"], check=True, capture_output=True)
    return json.loads((out / "results.json").read_text(encoding="utf-8")), rep


def test_docx_and_html_exist(built):
    _, rep = built
    assert (rep / "CT363_Smart_EV_Report.docx").stat().st_size > 50_000
    assert (rep / "CT363_Smart_EV_Report.html").stat().st_size > 10_000


def test_report_has_every_required_section(built):
    _, rep = built
    doc = Document(rep / "CT363_Smart_EV_Report.docx")
    heads = [p.text for p in doc.paragraphs if p.style.name == "Heading 1"]
    for needed in ["Abstract", "1. Introduction", "2. Literature Review", "3. Problem Formulation",
                   "4. System Architecture", "5. Methodology and Pseudocode", "6. Correctness",
                   "7. Complexity Analysis", "8. Implementation and Dataset", "9. Results",
                   "10. Conclusion and Limitations", "References"]:
        assert needed in heads, needed


def test_report_numbers_come_from_results(built):
    res, rep = built
    text = (rep / "CT363_Smart_EV_Report.html").read_text(encoding="utf-8")
    e1 = res["e1"]
    assert f"total cost {e1['cost']}" in text
    assert f"total benefit of {e1['allocation']['benefit']}" in text
    assert f"{res['verification']['agree']} of {res['verification']['instances']}" in text
    assert "Department of Computer Science" in text and "CT-24268" in text


def test_report_has_figures_and_pseudocode(built):
    _, rep = built
    doc = Document(rep / "CT363_Smart_EV_Report.docx")
    assert len(doc.inline_shapes) >= 8                     # architecture + 7 result figures
    body = "\n".join(c.text for t in doc.tables for row in t.rows for c in row.cells)
    for algo in ["KRUSKAL-MSF", "REPAIR", "MULTI-KNAPSACK-DP", "BRUTE-FORCE", "FIND(x)"]:
        assert algo in body, algo
